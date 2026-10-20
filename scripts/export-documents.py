"""Rebuild private blank templates. Development only: python-docx and reportlab."""
import json,re,zipfile
from pathlib import Path
from xml.sax.saxutils import escape
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
CAT=json.loads((ROOT/'documents/catalog.json').read_text())
OUT=ROOT/'documents/generated';OUT.mkdir(exist_ok=True)
FONT=Path('/usr/share/fonts/truetype/dejavu')
pdfmetrics.registerFont(TTFont('Dawkoz',str(FONT/'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('DawkozBold',str(FONT/'DejaVuSans-Bold.ttf')))
pdfmetrics.registerFontFamily('Dawkoz',normal='Dawkoz',bold='DawkozBold',italic='Dawkoz',boldItalic='DawkozBold')
styles={
 'title':ParagraphStyle('title',fontName='DawkozBold',fontSize=20,leading=25,textColor=colors.HexColor('#15273d'),spaceAfter=12),
 'meta':ParagraphStyle('meta',fontName='Dawkoz',fontSize=8,leading=12,textColor=colors.HexColor('#65758a'),spaceAfter=12),
 'note':ParagraphStyle('note',fontName='Dawkoz',fontSize=8,leading=12,textColor=colors.HexColor('#536478'),backColor=colors.HexColor('#edf2f8'),borderPadding=8,spaceAfter=16),
 'head':ParagraphStyle('head',fontName='DawkozBold',fontSize=10.5,leading=14,spaceBefore=10,spaceAfter=6,keepWithNext=True,textColor=colors.HexColor('#15273d')),
 'body':ParagraphStyle('body',fontName='Dawkoz',fontSize=9.2,leading=14,spaceAfter=9,textColor=colors.HexColor('#26364b'),allowWidows=0,allowOrphans=0),
}
def fill(text):return re.sub(r'\{\{([a-z_]+)\}\}',lambda m:CAT['fields'][m[1]]['default'] or '['+CAT['fields'][m[1]]['label']+']',text)
def p(text,style):return Paragraph(escape(fill(text)).replace('\n','<br/>'),styles[style])
def chrome(c,d):
 w,h=A4;c.saveState();c.setFont('DawkozBold',10);c.setFillColor(colors.HexColor('#15273d'));c.drawString(52,h-38,'DAWKOZ / DOKUMENTY')
 c.setStrokeColor(colors.HexColor('#dce4ee'));c.line(52,h-47,w-52,h-47);c.line(52,44,w-52,44)
 c.setFont('Dawkoz',7);c.setFillColor(colors.HexColor('#65758a'));c.drawString(52,31,'WZÓR B2B • PROJEKT DO UZGODNIENIA • WERSJA 1.0');c.drawRightString(w-52,31,str(d.page));c.restoreState()
for index,t in enumerate(CAT['templates'],1):
 story=[p(t['title'],'title'),p('Projekt do uzgodnienia • wzór B2B • wersja '+t['version'],'meta'),p(t['notes'] or CAT['notice'],'note')]
 for section in t['sections']:
  story.append(p(section['title'],'head'))
  story.extend(p(s,'body') for s in section['paragraphs'])
 SimpleDocTemplate(str(OUT/(t['id']+'.pdf')),pagesize=A4,leftMargin=46,rightMargin=46,topMargin=66,bottomMargin=61,title=t['title'],author='DAWKOZ').build(story,onFirstPage=chrome,onLaterPages=chrome)
 doc=Document();s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8);s.top_margin=Cm(2.1);s.bottom_margin=Cm(1.8)
 normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(7)
 for name,size in [('Title',22),('Heading 1',12)]:doc.styles[name].font.size=Pt(size);doc.styles[name].font.color.rgb=RGBColor.from_string('15273D')
 s.header.paragraphs[0].text='DAWKOZ / DOKUMENTY'
 foot=s.footer.paragraphs[0];foot.text='WZÓR B2B • PROJEKT DO UZGODNIENIA • WERSJA 1.0  |  '
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
 doc.add_paragraph(t['title'],'Title');doc.add_paragraph('Projekt do uzgodnienia • wersja '+t['version']);doc.add_paragraph(t['notes'] or CAT['notice'])
 md=['# '+t['title'],'','Projekt do uzgodnienia • wzór B2B • wersja '+t['version'],'',t['notes'] or CAT['notice'],'']
 for section in t['sections']:
  doc.add_paragraph(fill(section['title']),'Heading 1');md.extend(['## '+fill(section['title']),''])
  for text in section['paragraphs']:doc.add_paragraph(fill(text));md.extend([fill(text),''])
 doc.core_properties.title=t['title'];doc.core_properties.author='DAWKOZ';doc.save(OUT/(t['id']+'.docx'))
 (OUT/(t['id']+'.md')).write_text('\n'.join(md))
 print(index,t['id'])
readme='DAWKOZ — WZORY DOKUMENTÓW\n\n12 polskich wzorów B2B w formatach PDF, DOCX i Markdown.\nUzupełnij pola w nawiasach przed podpisaniem. Kwota 899 zł nie określa automatycznie statusu VAT, limitów ani czasów reakcji.\nPotwierdź wykonawcę i model współpracy z Kozbud. Wzory nie zostały sprawdzone przez prawnika dla konkretnej transakcji.\nWariant podstawowej umowy przewiduje licencję. Przeniesienie praw wymaga skoordynowania osobnego wzoru, łańcucha praw i formy podpisu.\nNie wpisuj haseł lub kluczy API do dokumentów.\n\n'
readme+='\n'.join(f'{i:02}. {t["title"]} — {t["id"]}' for i,t in enumerate(CAT['templates'],1))+'\n'
(OUT/'CZYTAJ.txt').write_text(readme)
zip_path=OUT/'dawkoz-wzory-dokumentow.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
 for file in sorted(OUT.iterdir()):
  if file.suffix in ['.pdf','.docx','.md','.txt']:z.write(file,file.name)
print('ZIP',zip_path.stat().st_size)
