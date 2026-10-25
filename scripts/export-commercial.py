"""Development export: commercial decision pack, editable cost workbook and presets."""
import csv, json, sys, zipfile
from pathlib import Path
from xml.sax.saxutils import escape
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from docx import Document
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT.parent / 'exports/dawkoz-oferta-startowa'
OUT.mkdir(parents=True, exist_ok=True)
O = json.loads((ROOT/'commercial/offer.json').read_text())
CAT = json.loads((ROOT/'documents/catalog.json').read_text())
M = lambda cents: f'{cents/100:,.2f}'.replace(',', ' ').replace('.', ',') + ' zł'
def evaluate(p, multiplier=1):
    net=round(p['monthlyCents']/(1+O['costAssumptions']['vatStressPercent']/100))
    costs=round((p['minutes']*p['assumedMinuteCostCents']+p['smsParts']*O['costAssumptions']['smsPartCostCents'])*multiplier)+p['fixedMonthlyCostCents']+round(p['monthlyLaborHours']*O['costAssumptions']['laborHourCostCents'])+p['monthlyReserveCents']
    return net, costs, net-costs

font=Path('/usr/share/fonts/truetype/dejavu')
pdfmetrics.registerFont(TTFont('DawkozCommercial',str(font/'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('DawkozCommercialBold',str(font/'DejaVuSans-Bold.ttf')))
pdfmetrics.registerFontFamily('DawkozCommercial',normal='DawkozCommercial',bold='DawkozCommercialBold',italic='DawkozCommercial',boldItalic='DawkozCommercialBold')
styles={
 'title':ParagraphStyle('title',fontName='DawkozCommercialBold',fontSize=21,leading=27,spaceAfter=16,textColor=colors.HexColor('#14283b')),
 'head':ParagraphStyle('head',fontName='DawkozCommercialBold',fontSize=11,leading=16,spaceBefore=15,spaceAfter=7,keepWithNext=True,textColor=colors.HexColor('#14283b')),
 'body':ParagraphStyle('body',fontName='DawkozCommercial',fontSize=9,leading=14,spaceAfter=8),
 'note':ParagraphStyle('note',fontName='DawkozCommercial',fontSize=8,leading=12,spaceAfter=12,textColor=colors.HexColor('#556577')),
}
def para(s,kind='body'):return Paragraph(escape(s).replace('\n','<br/>'),styles[kind])
def footer(c,d):
 c.saveState();c.setFont('DawkozCommercial',7);c.setFillColor(colors.HexColor('#65758a'));c.drawString(45,29,'DAWKOZ • WARUNKI STARTOWE • WERSJA '+O['version']);c.drawRightString(A4[0]-45,29,str(d.page));c.restoreState()
def write(name,title,note,sections):
 doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=sec.right_margin=Cm(1.8)
 doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
 doc.core_properties.title=title;doc.core_properties.author='DAWKOZ'
 doc.add_heading(title,0);doc.add_paragraph(note)
 story=[para(title,'title'),para(note,'note')];md=['# '+title,'',note,'']
 for head, paragraphs in sections:
  doc.add_heading(head,1);story.append(para(head,'head'));md+=['## '+head,'']
  for text in paragraphs:doc.add_paragraph(text);story.append(para(text));md += [text,'']
 doc.save(OUT/(name+'.docx'))
 SimpleDocTemplate(str(OUT/(name+'.pdf')),pagesize=A4,leftMargin=45,rightMargin=45,topMargin=42,bottomMargin=49,title=title,author='DAWKOZ').build(story,onFirstPage=footer,onLaterPages=footer)
 (OUT/(name+'.md')).write_text('\n'.join(md))

offer_sections=[]
for id,p in O['plans'].items():
 offer_sections.append((p['name'],[
  f'Wdrożenie od {M(p["setupCents"])}. Abonament {M(p["monthlyCents"])}/mies. Pakiet: {p["minutes"]} minut, {p["smsParts"]} części SMS oraz {p["changeMinutes"]} minut drobnych zmian miesięcznie.',
  next(x for x in CAT['presets'] if x['id']==id)['values']['scope'],
  next(x for x in CAT['presets'] if x['id']==id)['values']['deliverables'],
  f'Zakładany termin realizacji podstawowego zakresu: {p["deliveryBusinessDays"]} dni roboczych od otrzymania kompletnych materiałów, dostępów i zaliczki; termin potwierdzany w indywidualnej ofercie.'
 ]))
offer_sections += [
 ('Zakres i ceny końcowe',[
  'Wszystkie kwoty są cenami końcowymi do zapłaty; VAT, jeżeli dotyczy, mieści się w cenie. Faktyczny wykonawca, jego status i stawka VAT muszą być wskazane w umowie i dokumentach rozliczeniowych.',
  f'Warunki startowe wdrożeń AI obejmują 3 nowe projekty. Dostępność potwierdzana indywidualnie. Podstawowy zakres to jeden proces z wariantami rozmowy, baza do {O["plans"]["bot"]["faqAnswers"]} uzgodnionych odpowiedzi związanych z procesem na materiałach klienta i jedna standardowa integracja biznesowa, np. przekazanie podsumowania zgłoszenia do Telegrama, e-mail lub arkusza. Telefonia, AI i SMS są elementami pakietu, a nie dodatkowymi integracjami biznesowymi. Konto w panelu klienta wyceniane osobno.',
  'Dwie rundy uwag do scenariusza lub menu. Usuwanie niezgodności z zakresem nie jest ograniczone tymi rundami. Dodatkowe scenariusze, nowe API, numery premium, kampanie wychodzące i połączenia międzynarodowe wymagają osobnej wyceny.'
 ]),
 ('Zużycie, zgoda na dopłaty i pomiar',[
  'Powiadomienia przy 80% i 100% wykorzystania pakietu. Domyślny dodatkowy budżet: 0,00 zł. Po wyczerpaniu limitu nowe płatne rozmowy i SMS są wstrzymywane do akceptacji dodatkowego budżetu. Nie uruchamia się nieuzgodnionych płatnych przekierowań. Raport zużycia dostępny dla klienta.',
  'Każde połączenie i każdy płatny odcinek przekierowania zaokrąglane osobno w górę do pełnej minuty. Przy przekierowaniu mogą być naliczane dwa odcinki jednocześnie. SMS liczone w rozliczeniowych częściach zależnych od kodowania i długości. Niewykorzystane limity nie przechodzą na następny miesiąc.',
  'Domyślne maksima rozmowy: 5 minut AI, 15 minut przekierowania IVR. Ostatnia przyjęta rozmowa kończona w tym limicie nie powoduje nieuzgodnionej dopłaty.',
  f'Po akceptacji dodatkowego budżetu: IVR {M(O["billing"]["ivrExtraMinuteCents"])}/min, AI {M(O["billing"]["aiExtraMinuteCents"])}/min, SMS {M(O["billing"]["extraSmsPartCents"])}/część. Dodatkowe prace {M(O["billing"]["extraWorkHourCents"])}/godz. po akceptacji wyceny.'
 ]),
 ('Opieka i płatności',[
  'Wsparcie pon.–pt. 9:00–17:00 czasu polskiego, poza polskimi świętami. Rozpoczęcie obsługi awarii krytycznej do 8 godzin roboczych, pozostałych zgłoszeń do 2 dni roboczych. Czas reakcji nie oznacza terminu naprawy; postęp i termin przywrócenia działania komunikowane odrębnie.',
  '50% zaliczki po zawarciu umowy, 50% po wyraźnym odbiorze. Abonament z góry, od odbioru i aktywacji. Termin zapłaty faktury: 7 dni. Umowa utrzymania na czas nieoznaczony, wypowiedzenie 30 dni wcześniej ze skutkiem na koniec miesięcznego okresu rozliczeniowego.',
  'Zakończenie oferty startowej nie zmienia automatycznie ceny podpisanych umów. Zmiany ceny i zakresu wymagają uzgodnienia stron.'
 ]),
 ('Strony i automatyzacje',[
  'Strona firmowa od 1990,00 zł: jedna strona do 6 sekcji w jednym języku, z materiałów klienta. Automatyzacja od 990,00 zł: jeden proces łączący do 2 aplikacji. Dodatkowe podstrony, CMS, języki i płatne licencje wyceniane osobno.',
  'Opcjonalna opieka WWW od 149,00 zł/mies.: monitoring, kopia konfiguracji i 15 minut drobnych zmian. Opieka nad automatyzacją od 399,00 zł/mies.: monitoring uzgodnionego przepływu i 30 minut drobnych zmian. Domeny, hosting WWW i licencje klienta poza tymi abonamentami, rozliczane osobno. Opieka telefoniczna jest zawarta w pakietach telefonicznych.'
 ]),
 ('Konkretne wdrożenia automatyzacji',[
  f'Zapytania bez przepisywania od {M(O["otherServices"]["automation"]["examples"]["inquiries"]["setupCents"])}: jedno źródło, jedno miejsce docelowe, jeden proces łączący do dwóch aplikacji, np. formularz do arkusza.',
  f'Przypomnienia o terminach od {M(O["otherServices"]["automation"]["examples"]["reminders"]["setupCents"])}: jedno źródło terminów, jeden kanał e-mail albo SMS, jedna reguła wysyłki i uzgodniony szablon. Licencje i zużycie wiadomości osobno.',
  f'Faktury do sprawdzenia od {M(O["otherServices"]["automation"]["examples"]["invoices"]["setupCents"])}: jedna skrzynka, jeden uzgodniony typ tekstowego PDF, do pięciu pól przekazanych do arkusza do weryfikacji przez człowieka. Bez automatycznego księgowania, KSeF i OCR skanów; rozbudowy wyceniane osobno.',
  'Panel klienta: zakres i cena indywidualnie. Publiczne demo na dawkoz.com/panel-klienta.html pokazuje wyłącznie fikcyjne dane. Rzeczywiste konta, uprawnienia, baza, integracje i utrzymanie wymagają osobnego wdrożenia. Wersja przeglądarkowa unika potrzeby odrębnej aplikacji sklepowej; infrastruktura nadal ma koszt.'
 ])]
write('01-oferta-startowa','DAWKOZ — oferta startowa','Warunki do indywidualnej oferty B2B. Przed zawarciem umowy potwierdź dane stron, dostępność pakietu, zakres, dostawców i rzeczywisty status VAT.',offer_sections)

internal=[
 ('Pozycjonowanie cenowe i porównanie ofert',[
  'DAWKOZ nie powinien komunikować najniższej ceny. Publiczny cennik CallPage (https://www.callpage.pl/cennik) podaje podstawowego Agenta Głosowego AI za 399 zł netto miesięcznie przy subskrypcji rocznej, czyli 490,77 zł z VAT 23%. Strona https://www.callpage.pl/ai-agent deklaruje bezpłatne wdrożenie. Nie ustalono publicznego limitu minut AI i całego porównywalnego zakresu. Nie należy utożsamiać limitów callback z minutami bota.',
  f'Podstawowy poziom CallPage daje 5889,24 zł brutto za 12 miesięcy przed dodatkami. DAWKOZ Bot z wdrożeniem kosztuje {M(O["plans"]["bot"]["setupCents"]+12*O["plans"]["bot"]["monthlyCents"])} w pierwszych 12 miesiącach. Są to różne warunki i niezweryfikowany porównywalny zakres, więc nie jest to ranking opłacalności. Cena DAWKOZ wymaga uzasadnienia indywidualnym procesem, integracją, zakresem opieki i okresem wypowiedzenia, nie samą obietnicą AI.',
  'Datera (https://www.datera.pl/cennik/) publikuje 150/220/320/620 zł netto miesięcznie przy umowie na 24 miesiące. To pakiety centrali i komunikacji, a nie potwierdzony odpowiednik konwersacyjnego bota AI. Nie porównywać samych liczb minut bez funkcji, naliczania, wdrożenia i okresu umowy.',
  'Zwiększono Bot AI ze 150 do 200 minut bez podwyżki abonamentu, pozostawiono zestaw 250 minut. Podwojenie zestawu do 500 minut przy progu kosztu 1,20 zł/min, SMS, VAT 23%, pracy i rezerwie dawałoby stratę 124,11 zł miesięcznie. Faktyczne taryfy operatorów nie zostały odczytane; próg kosztu trzeba zmierzyć.'
 ]),
 ('Decyzja o rynku i pierwszej usłudze',[
  'Pierwszy kierunek: zarządcy nieruchomości. Wspólny scenariusz: przyjęcie zgłoszenia usterki, zebranie adresu i opisu, przekazanie sprawy do jednego narzędzia. Po pierwszych trzech wdrożeniach ocenić rozszerzenie na firmy serwisowe.',
  'Materiały sprzedażowe: nagranie 60–90 sekund rzeczywistej demonstracji, jednostronicowa oferta, opis zakresu i limitów, przykład raportu. Dane demonstracyjne oznaczać jako przykładowe. Publikować wyłącznie potwierdzone wyniki oraz materiały, do których jest uprawnienie.'
 ]),
 ('Architektura i dostawcy — plan do testów',[
  'Dla nowego prostego wdrożenia: telefonia i SMS w jednym narzędziu, preferencyjnie Twilio, oraz OpenAI w potoku transkrypcja → model tekstowy → synteza mowy. Kandydaci do testów: gpt-4o-mini-transcribe, gpt-4.1-mini, gpt-4o-mini-tts. Dostępność modeli, numerów i warunki kont trzeba potwierdzić przed ofertą.',
  'Na początku nie dodawać SMSAPI wyłącznie dla różnicy ceny przy 50–100 częściach SMS. Zmiana dostawcy ma sens po policzeniu: miesięczny wolumen × różnica kosztu części SMS wobec czasu integracji i utrzymania, albo gdy obecny dostawca przekracza ustalony próg kosztu.',
  'To decyzja projektowa do weryfikacji, nie potwierdzenie działającej integracji. Repozytorium strony nie zawiera silnika bota, połączeń z tymi API ani licznika wykorzystania.'
 ]),
 ('Model kosztów i rezerwa',[
  'Oficjalnych aktualnych stawek dostawców nie udało się odczytać w tej sesji. Wartości w arkuszu są wewnętrznymi progami do planowania, nie znalezionymi taryfami. Obliczenia uwzględniają pełne limity, VAT 23% po stronie sprzedaży, koszt pracy 100 zł/godz., koszty stałe i rezerwę. Nie zakładają odliczenia VAT od kosztów.',
  'Progi kosztu: minuta IVR do 0,60 zł, minuta AI z telefonią do 1,20 zł, część SMS do 0,30 zł — razem z opłatami operatora, podatkami i przeliczeniem walut. Miesięczny numer i hosting: budżet 35 zł IVR / 50 zł AI. Robocze koszty pracy i rezerwy są w arkuszu.',
  *[f'{p["name"]}: przychód po VAT {M(evaluate(p)[0])}, planowane koszty wraz z pracą i rezerwą {M(evaluate(p)[1])}, nadwyżka {M(evaluate(p)[2])}. Przy kosztach zmiennych wyższych o 50%: {M(evaluate(p,1.5)[2])} nadwyżki.' for p in O['plans'].values()],
  'Nadwyżka nie jest zyskiem do wypłaty: nie obejmuje PIT/CIT, ZUS, księgowości, pozyskania klienta i wszystkich kosztów ogólnych. Przekroczenie czasu pracy, stawek lub zakresu może odwrócić wynik.'
 ]),
 ('Panel i nowe automatyzacje',[
  'Publiczne demo jest materiałem do pokazania rozwiązania, nie repozytorium klientów. Dane są przykładowe, bez wgrywania plików i bez serwerowego zapisu. Realny panel wymaga uwierzytelnienia, izolacji firm, trwałej bazy, kopii z testem odtworzenia, retencji i audytu działań.',
  'Początek może korzystać z tego samego hostingu i interfejsu WWW; nie zamawiano nowych płatnych usług. Nie oznacza to bezkosztowego produkcyjnego przechowywania danych. PWA można dołożyć po ustaleniu funkcji i bezpieczeństwa. Biblioteka /app pozostaje osobnym narzędziem operatora.',
  'Sprzedawać jeden konkretny proces, a nie ogólne „AI w księgowości”. Podstawowe faktury: tekstowy PDF uzgodnionego typu, do pięciu pól, ręczna weryfikacja. OCR dowolnych skanów, autoksięgowanie i KSeF wymagają innego zakresu i wyceny. Koszty licencji i zużycia oraz zasady obsługi błędów przed umową.'
 ]),
 ('Warunki przyjęcia pierwszego klienta',[
  'Przed podpisaniem oferty wykonać minimum 20 zróżnicowanych rozmów testowych na docelowym stosie: adres znany/nieznany, przerwanie rozmowy, brak danych, błąd integracji i przekazanie sprawy człowiekowi. Sprawdzić rzeczywiste koszty na zestawieniu dostawców, nie tylko oszacowanie modelu.',
  'Potwierdzić koszt numeru i dostępność SMS, opłaty za oba odcinki przekierowania, naliczanie rozpoczętej minuty i segmentów SMS. Przy przekroczeniu progu kosztu zmienić technologię lub przygotować inną ofertę przed podpisaniem. Nie zmieniać już uzgodnionej ceny bez akceptacji klienta.',
  'Przed aktywacją skonfigurować i przetestować: licznik miesięczny, reset okresu, raport, alerty 80% i 100%, maksymalny czas rozmowy, ograniczenie wydatków u dostawców oraz blokadę nowych płatnych zdarzeń bez zgody. Podpisany dokument i zapis na stronie nie zastępują tych zabezpieczeń w systemie bota.',
  'Przetestować kopię konfiguracji i odtworzenie; ustalić faktyczny wykaz dostawców, regiony, retencję oraz zasady dostępu do danych. Informację o AI i ewentualnym nagrywaniu dopasować do rzeczywistych funkcji.'
 ]),
 ('Przebieg zlecenia i dokumenty',[
  '1. Brief: proces, miesięczny ruch, długość rozmów, integracja, osoba kontaktowa, obecny operator i dane do oferty.',
  '2. Weryfikacja kosztów i wykonalności. Oferta z nazwą pakietu, limitem, ceną końcową, zakresem i potwierdzoną dostępnością.',
  '3. Umowa wdrożeniowa i zakres, umowa opieki, w razie potrzeby powierzenie danych. Ustalenie uprawnionego wykonawcy i reprezentacji. 50% zaliczki po zawarciu umowy.',
  '4. Konfiguracja w podstawowym zakresie, dwie rundy uwag do scenariusza/menu i 5 dni roboczych na testy klienta. Usuwanie niezgodności z zakresem osobno od rozbudowy.',
  '5. Wyraźny protokół odbioru, płatność pozostałej części, przekazanie instrukcji i kont. Abonament od odbioru i aktywacji.',
  '6. Po pierwszym miesiącu przegląd kosztu, zużycia, usterek i czasu pracy. Następnie raport miesięczny i monitorowanie warunków opieki.'
 ]),
 ('Wykonawca, Kozbud i podatki — fakty do potwierdzenia',[
  'Dawid Kozula: koordynacja techniczna i rozwój marki DAWKOZ. NIP 6272742701 został wskazany jako NIP Kozbud. Pełna nazwa prawna, właściciel, adres, reprezentacja i status VAT nie zostały niezależnie potwierdzone.',
  'Możliwy model: Kozbud jako rzeczywisty wykonawca umowy i sprzedawca, a Dawid jako realizator na uzgodnionej podstawie. Wymaga to rzeczywistej zgody właściciela i ustalenia rozliczeń z księgowością. Tych faktów i cudzej zgody nie można ustanowić samym dokumentem.',
  'Do potwierdzenia pozostają pola wykonawcy, reprezentacji, rachunku i statusu VAT. Ceny końcowe już obejmują VAT, jeśli jest należny; nie dopisywać go klientowi ponad uzgodnioną cenę. Nie łączyć NIP Kozbud z inną osobą jako sprzedawcą.'
 ]),
 ('Ewidencja trzech wdrożeń i dalsze ceny',[
  'Dostępność warunków startowych sprawdzać przy każdej ofercie. Prowadzić prosty rejestr: numer oferty, pakiet, data ważności, status podpisania, otrzymana zaliczka i zużyte miejsce. Nie publikować fikcyjnego licznika pozostałych miejsc.',
  'Po trzech nowych wdrożeniach AI porównać faktyczny czas realizacji i marżę. Dopiero na tej podstawie ustalić ceny nowych ofert. Podpisane umowy zachowują uzgodnione warunki; zmiany wymagają akceptacji stron.'
 ]),
 ('Plan pierwszych dwóch tygodni',[
  'Dni 1–3: komplet danych sprzedawcy, testy dostawców i kosztów, dopięcie demonstracji oraz materiałów do oferty.',
  'Dni 4–7: uzupełnienie profilu LinkedIn, publikacja demonstracji i jednego konkretnego opisu rozwiązania. Rozmowy z osobami, które zgłosiły zainteresowanie lub polecono kontakt.',
  'Dni 8–14: kwalifikacja zainteresowanych, briefy, indywidualne oferty i prezentacje. Ewidencjonować pytania, obiekcje i rzeczywisty zakres, bez obiecywania wyników sprzedażowych.'
 ])]
write('02-plan-startu','DAWKOZ — plan startu i kalkulacja','Dokument wewnętrzny. Założenia kosztowe nie są zweryfikowanymi stawkami dostawców. Oferta wymaga weryfikacji wykonalności przed przyjęciem zlecenia.',internal)

wb=Workbook();a=wb.active;a.title='Założenia'
a.append(['DAWKOZ — założenia kosztowe','Wartość'])
a.append(['NIEZWERYFIKOWANE PROGI PLANOWANIA. Zastąp rzeczywistymi kosztami przed przyjęciem zlecenia.'])
a.append(['Pozycja','Wartość'])
a.append(['VAT sprzedaży — scenariusz ostrożny',O['costAssumptions']['vatStressPercent']/100])
a.append(['Koszt godziny pracy (zł)',O['costAssumptions']['laborHourCostCents']/100])
a.append(['Koszt części SMS (zł)',O['costAssumptions']['smsPartCostCents']/100])
a.append(['Mnożnik stresowy kosztów zmiennych',O['costAssumptions']['stressMultiplierPercent']/100])
a['B4'].number_format='0%';a['B7'].number_format='0%'
a.append(['Nadwyżka jest przed PIT/CIT, ZUS, księgowością, sprzedażą i innymi kosztami ogólnymi.'])
monthly=wb.create_sheet('Abonamenty')
monthly.append(['Pakiet','Cena końcowa','Przychód po VAT','Minuty','Części SMS','Koszt minuty','Koszty zmienne','Numer i hosting','Praca: godziny','Koszt pracy','Rezerwa','Koszty razem','Nadwyżka','Nadwyżka w stresie'])
setup=wb.create_sheet('Wdrożenia');setup.append(['Pakiet','Cena końcowa','Przychód po VAT','Godziny pracy','Koszt pracy','Koszty zewnętrzne','Rezerwa','Koszty razem','Nadwyżka','Pierwszy rok — klient'])
for row,(id,p) in enumerate(O['plans'].items(),2):
 monthly.append([p['name'],p['monthlyCents']/100,f'=ROUND(B{row}/(1+\'Założenia\'!$B$4),2)',p['minutes'],p['smsParts'],p['assumedMinuteCostCents']/100,f'=ROUND(D{row}*F{row}+E{row}*\'Założenia\'!$B$6,2)',p['fixedMonthlyCostCents']/100,p['monthlyLaborHours'],f'=ROUND(I{row}*\'Założenia\'!$B$5,2)',p['monthlyReserveCents']/100,f'=SUM(G{row}:H{row})+J{row}+K{row}',f'=C{row}-L{row}',f'=C{row}-(G{row}*\'Założenia\'!$B$7+H{row}+J{row}+K{row})'])
 setup.append([p['name'],p['setupCents']/100,f'=ROUND(B{row}/(1+\'Założenia\'!$B$4),2)',p['setupLaborHours'],f'=D{row}*\'Założenia\'!$B$5',p['setupExternalCostCents']/100,p['setupReserveCents']/100,f'=SUM(E{row}:G{row})',f'=C{row}-H{row}',f'=B{row}+12*Abonamenty!B{row}'])
register=wb.create_sheet('Oferty startowe')
register.append(['Miejsce','Numer oferty','Pakiet','Ważność oferty','Data podpisania','Zaliczka otrzymana','Status'])
for slot in range(1,O['launch']['aiProjects']+1):register.append([slot,'','','','','',''])
automations=wb.create_sheet('Automatyzacje')
automations.append(['Zakres startowy','Cena końcowa od','Przychód po VAT','Godziny pracy','Koszt pracy','Koszty zewnętrzne','Rezerwa','Koszty razem','Nadwyżka'])
names={'inquiries':'Zapytania do jednego narzędzia','reminders':'Przypomnienia: jeden kanał','invoices':'Tekstowy PDF do weryfikacji'}
for row,(id,p) in enumerate(O['otherServices']['automation']['examples'].items(),2):
 automations.append([names[id],p['setupCents']/100,f'=ROUND(B{row}/(1+\'Założenia\'!$B$4),2)',p['setupLaborHours'],f'=D{row}*\'Założenia\'!$B$5',p['externalCostCents']/100,p['reserveCents']/100,f'=SUM(E{row}:G{row})',f'=C{row}-H{row}'])
for ws in wb:
 ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='15273D');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=42
 for col in range(1,ws.max_column+1):ws.column_dimensions[__import__('openpyxl').utils.get_column_letter(col)].width=19 if col>1 else 31
 for cells in ws.iter_rows(min_row=2):
  for c in cells:
   c.alignment=Alignment(vertical='top',wrap_text=True)
   if c.column>1 and c.value is not None:
    c.fill=PatternFill('solid',fgColor='EDF4FF' if c.data_type!='f' else 'F1F6EF')
    if isinstance(c.value,(int,float)) or c.data_type=='f':c.number_format='#,##0.00'
 ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
a.column_dimensions['A'].width=80;a['B4'].number_format='0%';a['B7'].number_format='0%'
a.row_dimensions[2].height=48;a.row_dimensions[8].height=48
wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True)
wb.save(OUT/'03-kalkulacja-kosztow.xlsx')
with (OUT/'03-kalkulacja-podsumowanie.csv').open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.writer(f,delimiter=';');writer.writerow(['Pakiet','Cena końcowa/mies.','Przychód po VAT 23%','Planowane koszty z pracą i rezerwą','Nadwyżka przed PIT/CIT i ZUS','Nadwyżka przy wzroście kosztów zmiennych o 50%'])
 for p in O['plans'].values():
  net,cost,surplus=evaluate(p);writer.writerow([p['name'],M(p['monthlyCents']),M(net),M(cost),M(surplus),M(evaluate(p,1.5)[2])])

for preset in CAT['presets']:
 values={key:f.get('default','') for key,f in CAT['fields'].items()};values.update(preset['values'])
 project={'type':'DAWKOZ_DOCUMENT_PROJECT','version':1,'templateId':'oferta','values':values}
 (OUT/('projekt-'+preset['id']+'.json')).write_text(json.dumps(project,ensure_ascii=False,indent=2)+'\n')
 # Render a ready-to-complete offer without inventing client or seller identity.
 template=next(t for t in CAT['templates'] if t['id']=='oferta')
 def fill(s):
  import re
  return re.sub(r'\{\{([a-z_]+)\}\}',lambda m:values.get(m[1]) or '['+CAT['fields'][m[1]]['label']+']',s)
 sections=[(fill(s['title']),[fill(p) for p in s['paragraphs']]) for s in template['sections']]
 write('oferta-'+preset['id'],'Oferta — '+O['plans'][preset['id']]['name'],'Projekt do uzupełnienia: dane stron, status VAT, termin i konkretny odbiorca integracji wymagają potwierdzenia przed wysłaniem klientowi.',sections)

for name in ['audyt-oferty.md','audyt-prawny.md']:
 source=ROOT/'docs'/name
 if source.exists():(OUT/name).write_text(source.read_text())
(OUT/'CZYTAJ.txt').write_text('DAWKOZ — GOTOWA OFERTA STARTOWA\n\n01: warunki dla klientów. 02: plan wewnętrzny z progami kosztów, procesem i zadaniami. 03: edytowalny Excel i podsumowanie CSV. Audyty: ustalenia o ofercie i prawie.\nOferty IVR, bot i combo: PDF/DOCX/MD do uzupełnienia danymi stron. Pliki projekt-*.json można wczytać w panelu dokumentów dawkoz.com/app.\nNiebieskie pola Excela to założenia, zielone to formuły. Excel przelicza je po otwarciu. Stawki dostawców nie zostały zweryfikowane; to model planowania. Arkusz obejmuje też trzy konkretne automatyzacje.\nPotwierdź rzeczywistego sprzedawcę, status VAT i koszty dostawców. Polityka pozostaje projektem do uzupełnienia rzeczywistymi danymi. Limity na stronie i w dokumentach nie są wdrożeniem ograniczeń w silniku bota. Demo panelu nie jest kontem klienta.\n')
archive=OUT.parent/'dawkoz-oferta-startowa.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
 for file in sorted(OUT.iterdir()):
  if file.is_file():z.write(file,file.name)
print(archive,archive.stat().st_size)
